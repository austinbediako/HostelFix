# Master Thesis Compiler for HostelFix 150+ Page Undergraduate Project Document
# Conforming strictly to Project Template - Undergraduate.docx and Thesis or Dissertation template.docx.pdf
# Candidate: Asante Gideon Kwadwo (11287773)
# Supervisor: Prof. Winfred Yaokumah
# Department of Computer Science, University of Ghana, Legon
# Date: March 2026

import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from data_prelims import PRELIMS_DATA
from data_ch1 import CH1_DATA
from data_ch2 import CH2_DATA
from data_ch3 import CH3_DATA
from data_ch4 import CH4_DATA
from data_ch5 import CH5_DATA
from data_references import REFERENCES_DATA
from data_appendices import APPENDICES_DATA

def sanitize_text(text):
    if not isinstance(text, str):
        return text
    text = text.replace('\u2014', ' - ')  # Em dash to hyphen
    text = text.replace('\u2013', '-')    # En dash to hyphen
    text = text.replace('--', ' - ')
    return text

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
if os.path.exists(os.path.join(SCRIPT_DIR, '..', 'docs')):
    REPO_ROOT = os.path.abspath(os.path.join(SCRIPT_DIR, '..'))
else:
    REPO_ROOT = SCRIPT_DIR

def compile_document():
    print("Initializing Master Document Compiler...")
    print(f"Project root resolved to: {REPO_ROOT}")
    doc = docx.Document()

    # Base Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal_style.paragraph_format.space_after = Pt(6)
    normal_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # XML Helper Functions
    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def add_page_number(run):
        fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        run._r.append(fldSimple)

    # Paragraph formatting helpers
    def p_title_center(text, size=14, bold=True, space_before=0, space_after=6):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        run = p.add_run(sanitize_text(text))
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
        run.font.bold = bold
        return p

    def p_body(text, space_after=6, bold_prefix=None, italic=False):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        if bold_prefix:
            r_pre = p.add_run(sanitize_text(bold_prefix))
            r_pre.font.name = 'Times New Roman'
            r_pre.font.size = Pt(12)
            r_pre.font.bold = True
        run = p.add_run(sanitize_text(text))
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.italic = italic
        return p

    def p_h1(text, space_before=12, space_after=6):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        run = p.add_run(sanitize_text(text))
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(26, 54, 93)
        return p

    def p_h2(text, space_before=10, space_after=4):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        run = p.add_run(sanitize_text(text))
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.bold = True
        return p

    def add_image_figure(img_path, caption, width_in=5.8):
        resolved_img = os.path.join(REPO_ROOT, img_path) if not os.path.isabs(img_path) else img_path
        if not os.path.exists(resolved_img):
            print(f"Warning: image path does not exist: {img_path}")
            return
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(resolved_img, width=Inches(width_in))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(10)
        run_cap = p_cap.add_run(sanitize_text(caption))
        run_cap.font.name = 'Times New Roman'
        run_cap.font.size = Pt(11)
        run_cap.font.bold = True
        run_cap.font.italic = True
        run_cap.font.color.rgb = RGBColor(43, 108, 176)

    def add_styled_table(caption, headers, rows, col_widths=None):
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_cap.paragraph_format.space_before = Pt(10)
        p_cap.paragraph_format.space_after = Pt(4)
        run_cap = p_cap.add_run(sanitize_text(caption))
        run_cap.font.name = 'Times New Roman'
        run_cap.font.size = Pt(11)
        run_cap.font.bold = True
        run_cap.font.color.rgb = RGBColor(26, 54, 93)

        num_cols = len(headers)
        table = doc.add_table(rows=len(rows) + 1, cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Header Row
        hdr_cells = table.rows[0].cells
        for col_idx, header_text in enumerate(headers):
            cell = hdr_cells[col_idx]
            cell.text = sanitize_text(header_text)
            set_cell_background(cell, "1A365D")
            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10)
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)

        # Body Rows
        for row_idx, row_data in enumerate(rows):
            row_cells = table.rows[row_idx + 1].cells
            fill_color = "F7FAFC" if row_idx % 2 == 1 else "FFFFFF"
            for col_idx, cell_value in enumerate(row_data):
                cell = row_cells[col_idx]
                cell.text = sanitize_text(str(cell_value))
                set_cell_background(cell, fill_color)
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(9.5)

        if col_widths and len(col_widths) == num_cols:
            for row in table.rows:
                for col_idx, width in enumerate(col_widths):
                    row.cells[col_idx].width = Inches(width)

        p_spacer = doc.add_paragraph()
        p_spacer.paragraph_format.space_after = Pt(6)

    def add_code_block(caption, code_text):
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_cap.paragraph_format.space_before = Pt(8)
        p_cap.paragraph_format.space_after = Pt(2)
        r_cap = p_cap.add_run(sanitize_text(caption))
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(11)
        r_cap.font.bold = True

        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "F8F9FA")
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        cell.width = Inches(6.0)

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(sanitize_text(code_text))
        r.font.name = 'Courier New'
        r.font.size = Pt(9.0)

        p_spacer = doc.add_paragraph()
        p_spacer.paragraph_format.space_after = Pt(6)

    # Configure Margins for Section 1 (Preliminary Pages)
    section1 = doc.sections[0]
    section1.top_margin = Inches(1.0)
    section1.bottom_margin = Inches(1.0)
    section1.left_margin = Inches(1.5)
    section1.right_margin = Inches(1.0)
    section1.different_first_page_header_footer = True

    # ==========================================
    # 1. COVER PAGE 1
    # ==========================================
    print("Generating Cover Page 1...")
    crest_path = os.path.join(REPO_ROOT, "docs/diagrams/ug_crest.png")
    if os.path.exists(crest_path):
        p_crest = doc.add_paragraph()
        p_crest.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_crest.paragraph_format.space_before = Pt(36)
        p_crest.paragraph_format.space_after = Pt(18)
        p_crest.add_run().add_picture(crest_path, width=Inches(1.8))

    p_title_center("HOSTELFIX: A ROLE-BASED MAINTENANCE REPORTING AND TRACKING SYSTEM FOR TERTIARY RESIDENTIAL FACILITIES", size=16, bold=True, space_before=12, space_after=6)
    p_title_center("A Case Study of University of Ghana Student Halls of Residence", size=13, bold=False, space_before=0, space_after=36)

    p_title_center("BY", size=12, bold=True, space_before=12, space_after=12)
    p_title_center("ASANTE GIDEON KWADWO", size=14, bold=True, space_before=0, space_after=4)
    p_title_center("(STUDENT ID: 11287773)", size=12, bold=False, space_before=0, space_after=48)

    p_title_center(
        "A Project Report submitted to the Department of Computer Science, "
        "College of Basic and Applied Sciences, University of Ghana, Legon, "
        "in partial fulfillment of the requirements for the award of\n"
        "Bachelor of Science in Computer Science Degree",
        size=12, bold=False, space_before=12, space_after=48
    )

    p_title_center("MARCH 2026", size=12, bold=True, space_before=24, space_after=0)
    doc.add_page_break()

    # ==========================================
    # 2. COVER PAGE 2 / DECLARATION
    # ==========================================
    print("Generating Declaration Page...")
    p_title_center("DECLARATION", size=14, bold=True, space_before=18, space_after=18)

    p_h2("STUDENT DECLARATION")
    p_body(
        "I, Asante Gideon Kwadwo, hereby declare that this project report entitled "
        "'HOSTELFIX: A ROLE-BASED MAINTENANCE REPORTING AND TRACKING SYSTEM FOR TERTIARY RESIDENTIAL FACILITIES' "
        "is the result of my own independent research work conducted under the supervision of Prof. Winfred Yaokumah "
        "in the Department of Computer Science, University of Ghana, Legon. This work has not been presented, either in "
        "whole or in part, for any other degree or diploma in this university or elsewhere. All literary sources, empirical data, "
        "and intellectual contributions from other authors have been duly acknowledged and referenced in accordance with academic standards."
    )

    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(30)
    p_sig.paragraph_format.space_after = Pt(4)
    p_sig.add_run("Signature: ___________________________               Date: ___________________________")
    p_body("Asante Gideon Kwadwo\n(Candidate - 11287773)", bold_prefix=None, space_after=24)

    p_h2("SUPERVISOR CERTIFICATION")
    p_body(
        "I hereby certify that the preparation and presentation of this final year project report were supervised by me "
        "in accordance with the guidelines on supervision of undergraduate projects laid down by the Department of Computer Science "
        "and the University of Ghana."
    )

    p_sig2 = doc.add_paragraph()
    p_sig2.paragraph_format.space_before = Pt(30)
    p_sig2.paragraph_format.space_after = Pt(4)
    p_sig2.add_run("Signature: ___________________________               Date: ___________________________")
    p_body("Prof. Winfred Yaokumah\n(Project Supervisor, Department of Computer Science)", bold_prefix=None, space_after=12)
    doc.add_page_break()

    # ==========================================
    # 3. ABSTRACT
    # ==========================================
    print("Generating Abstract...")
    p_title_center("ABSTRACT", size=14, bold=True, space_before=18, space_after=18)
    for p_text in PRELIMS_DATA["abstract"].split("\n\n"):
        p_body(p_text)
    p_body("Keywords: maintenance reporting, tertiary residential halls, role-based access control, MongoDB, Next.js, Express, workflow design, audit trail, information systems.", bold_prefix="Keywords: ", space_after=12)
    doc.add_page_break()

    # ==========================================
    # 4. DEDICATION
    # ==========================================
    print("Generating Dedication...")
    p_title_center("DEDICATION", size=14, bold=True, space_before=18, space_after=18)
    for p_text in PRELIMS_DATA["dedication"].split("\n\n"):
        p_body(p_text)
    doc.add_page_break()

    # ==========================================
    # 5. ACKNOWLEDGEMENTS
    # ==========================================
    print("Generating Acknowledgements...")
    p_title_center("ACKNOWLEDGEMENTS", size=14, bold=True, space_before=18, space_after=18)
    for p_text in PRELIMS_DATA["acknowledgements"].split("\n\n"):
        p_body(p_text)
    doc.add_page_break()

    # ==========================================
    # 6. TABLE OF CONTENTS
    # ==========================================
    print("Generating Table of Contents...")
    p_title_center("TABLE OF CONTENTS", size=14, bold=True, space_before=18, space_after=18)
    toc_items = [
        ("DECLARATION", "ii"),
        ("ABSTRACT", "iii"),
        ("DEDICATION", "iv"),
        ("ACKNOWLEDGEMENTS", "v"),
        ("TABLE OF CONTENTS", "vi"),
        ("LIST OF FIGURES", "viii"),
        ("LIST OF TABLES", "x"),
        ("LIST OF ABBREVIATIONS AND ACRONYMS", "xii"),
        ("CHAPTER ONE: INTRODUCTION", "1"),
        ("  1.1 Introduction and Theoretical Constructs", "1"),
        ("  1.2 Background to the Study", "4"),
        ("  1.3 Research Problem Statement", "7"),
        ("  1.4 Research Questions", "9"),
        ("  1.5 Research Aims and Objectives", "10"),
        ("  1.6 Limitations and Scope of the Study", "12"),
        ("  1.7 Research Methodology (Brief Overview)", "13"),
        ("  1.8 Organization of the Project Document", "14"),
        ("CHAPTER TWO: LITERATURE REVIEW", "16"),
        ("  2.1 Introduction and Theoretical Framework", "16"),
        ("  2.2 Maintenance Engineering Paradigms and Asset Lifecycles", "20"),
        ("  2.3 Empirical Review of Tertiary Facility Maintenance Systems", "24"),
        ("  2.4 Review and Critical Comparison of Similar Systems", "29"),
        ("  2.5 Technological Foundations and Architectural Paradigms", "36"),
        ("  2.6 Summary and Research Gaps Addressed", "42"),
        ("CHAPTER THREE: SYSTEM ANALYSIS AND DESIGN", "44"),
        ("  3.1 Software Development Methodology", "44"),
        ("  3.2 Requirements Engineering", "48"),
        ("  3.3 Input Design", "55"),
        ("  3.4 Output Design", "58"),
        ("  3.5 Database Design", "61"),
        ("  3.6 System Modeling and UML Diagrams", "72"),
        ("  3.7 Hardware and Software Specifications", "102"),
        ("CHAPTER FOUR: IMPLEMENTATION AND EVALUATION", "106"),
        ("  4.1 System Implementation Overview and Module Decomposition", "106"),
        ("  4.2 Comprehensive User Interface and Workflow Walkthrough", "109"),
        ("  4.3 Testing Strategy and Verification Architecture", "128"),
        ("  4.4 Test Cases, Test Data, and Execution Results", "131"),
        ("  4.5 Performance, Latency, and Scalability Evaluation", "142"),
        ("  4.6 Security Audit and Vulnerability Assessment", "146"),
        ("  4.7 User Acceptance Testing (UAT) and Evaluation Analysis", "150"),
        ("CHAPTER FIVE: CONCLUSION AND FUTURE WORKS", "156"),
        ("  5.1 Summary of Results", "156"),
        ("  5.2 Contributions to Knowledge and Facility Management", "159"),
        ("  5.3 Limitations of the Study", "161"),
        ("  5.4 Strategic Recommendations", "162"),
        ("  5.5 Future Work and Research Directions", "164"),
        ("  5.6 Concluding Remarks", "167"),
        ("REFERENCES", "169"),
        ("APPENDICES", "176"),
        ("  APPENDIX A: Complete REST API Specification", "176"),
        ("  APPENDIX B: Database Schemas and Mongoose Data Models", "182"),
        ("  APPENDIX C: Automated Test Suite Implementation Code", "188"),
        ("  APPENDIX D: System Deployment and Environment Configuration Guide", "191"),
        ("  APPENDIX E: User Acceptance Testing Evaluation Instrument (SUS)", "194"),
        ("  APPENDIX F: Audit Trail and Security Event Log Schema", "196")
    ]
    for title, page_str in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run(sanitize_text(title))
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(11)
        if title.startswith("CHAPTER") or title in ["DECLARATION", "ABSTRACT", "REFERENCES", "APPENDICES"]:
            r1.font.bold = True
        # Dotted leader
        dots_count = max(5, 75 - len(title))
        r_dots = p.add_run(" " + ". " * (dots_count // 2) + " ")
        r_dots.font.name = 'Times New Roman'
        r_dots.font.size = Pt(10)
        r_dots.font.color.rgb = RGBColor(160, 160, 160)
        r2 = p.add_run(page_str)
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(11)
        r2.font.bold = True
    doc.add_page_break()

    # ==========================================
    # 7. LIST OF FIGURES
    # ==========================================
    print("Generating List of Figures...")
    p_title_center("LIST OF FIGURES", size=14, bold=True, space_before=18, space_after=18)
    figures_list = [
        ("Figure 3.1: Multi-Tier Modular Monolith Architecture of the HostelFix Platform", "73"),
        ("Figure 3.2: UML Component Diagram Illustrating Internal Modular Boundaries", "75"),
        ("Figure 3.3: Physical Network Deployment and Infrastructure Topology Diagram", "77"),
        ("Figure 3.4: Operational State Machine Diagram Governed by ADR 001", "79"),
        ("Figure 3.5: Entity-Relationship Diagram (ERD) of the HostelFix Database", "82"),
        ("Figure 3.6: UML Domain Class Diagram Illustrating Model Classes and Associations", "85"),
        ("Figure 3.7: Comprehensive UML Use Case Diagram across Five Stakeholder Roles", "88"),
        ("Figure 3.8: UML Sequence Diagram: Dual-Token JWT Authentication and Refresh Cycle", "93"),
        ("Figure 3.9: UML Sequence Diagram: Complete Issue Reporting, Triage, and Resolution Lifecycle", "95"),
        ("Figure 3.10: Activity Flowchart: Student Incident Reporting and Validation Workflow", "97"),
        ("Figure 3.11: Activity Flowchart: Hall Manager Triage and Assignment Workflow", "99"),
        ("Figure 3.12: Activity Flowchart: Maintenance Technician Work Order Execution Workflow", "101"),
        ("Figure 4.1: HostelFix Authentication and Role-Based Login Interface", "110"),
        ("Figure 4.2: Student Resident Command Dashboard", "112"),
        ("Figure 4.3: Student Maintenance Issue Reporting Interface", "114"),
        ("Figure 4.4: Student Filterable Ticket History Queue", "116"),
        ("Figure 4.5: Issue Detail, Chronological Timeline, and Dispute Re-opening View", "118"),
        ("Figure 4.6: Hall Manager Operations Dashboard", "120"),
        ("Figure 4.7: Hall Manager Issue Management and Artisan Assignment Modal", "122"),
        ("Figure 4.8: Maintenance Personnel Mobile Work Order Dashboard", "124"),
        ("Figure 4.9: University Administrator Executive Cross-Hall Dashboard", "126"),
        ("Figure 4.10: University Administrator Analytics and CSV Export Console", "127"),
        ("Figure 4.11: System Administrator Governance Console and User Management", "128"),
        ("Figure 4.12: HostelFix Automated Testing Pyramid Architecture", "130"),
        ("Figure 4.13: API Endpoint Latency Benchmarks under Concurrent Load (100 Concurrency)", "144"),
        ("Figure 4.14: System Usability Scale (SUS) Score by Stakeholder Cohort", "152")
    ]
    for fig_title, page_str in figures_list:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run(sanitize_text(fig_title))
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(11)
        dots_count = max(5, 75 - len(fig_title))
        r_dots = p.add_run(" " + ". " * (dots_count // 2) + " ")
        r_dots.font.name = 'Times New Roman'
        r_dots.font.size = Pt(10)
        r_dots.font.color.rgb = RGBColor(160, 160, 160)
        r2 = p.add_run(page_str)
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(11)
        r2.font.bold = True
    doc.add_page_break()

    # ==========================================
    # 8. LIST OF TABLES
    # ==========================================
    print("Generating List of Tables...")
    p_title_center("LIST OF TABLES", size=14, bold=True, space_before=18, space_after=18)
    tables_list = [
        ("Table 2.1: Comprehensive Comparative Architectural Matrix of Maintenance Management Systems", "34"),
        ("Table 3.1: Comprehensive Functional Requirements Specification (FR-01 to FR-30)", "50"),
        ("Table 3.2: Non-Functional Requirements Specification Mapped to ISO/IEC 25010 Quality Model", "53"),
        ("Table 3.3: Detailed Input Data Specifications and Validation Rules", "56"),
        ("Table 3.4: Data Dictionary for 'users' Collection", "63"),
        ("Table 3.5: Data Dictionary for 'issues' Collection", "65"),
        ("Table 3.6: Data Dictionary for 'issue_events' Collection (Audit Trail)", "67"),
        ("Table 3.7: Data Dictionary for 'residences' Collection (Halls of Residence)", "68"),
        ("Table 3.8: Data Dictionary for 'refresh_tokens' Collection", "69"),
        ("Table 3.9: Data Dictionary for 'assignments' Collection", "70"),
        ("Table 3.10: Data Dictionary for 'notifications' Collection", "71"),
        ("Table 3.11: Data Dictionary for 'audit_logs' Collection", "71"),
        ("Table 3.12: Data Dictionary for 'system_settings' Collection", "72"),
        ("Table 3.13: Use Case Specification UC-01: Authenticate with Credentials", "89"),
        ("Table 3.14: Use Case Specification UC-02: Submit Maintenance Ticket", "90"),
        ("Table 3.15: Use Case Specification UC-03: View and Filter Personal Issue History", "91"),
        ("Table 3.16: Use Case Specification UC-04: Dispute Resolution (Reopen Ticket)", "92"),
        ("Table 3.17: Use Case Specification UC-05: Acknowledge Submitted Ticket (ADR 001)", "92"),
        ("Table 3.18: Use Case Specification UC-06: Reject Illegitimate Ticket with Reason", "93"),
        ("Table 3.19: Use Case Specification UC-07: Update Status to Resolved", "94"),
        ("Table 3.20: Use Case Specification UC-08: Execute Repair and Update Status to In Progress", "94"),
        ("Table 3.21: Use Case Specification UC-09: Complete Work Order and Submit Resolution Evidence", "95"),
        ("Table 3.22: Use Case Specification UC-10: View Cross-Hall Facility Analytics and MTTR", "96"),
        ("Table 3.23: Use Case Specification UC-11: Provision User Accounts and Assign Hall Roles", "96"),
        ("Table 3.24: Use Case Specification UC-12: Inspect System Audit Logs and Trigger Manual Escalation", "97"),
        ("Table 3.25: Comprehensive Hardware and Software Environment Specifications", "104"),
        ("Table 4.1: Authentication and Session Security Test Suite (TC-AUTH-01 to TC-AUTH-08)", "132"),
        ("Table 4.2: Student Issue Reporting and Validation Test Suite (TC-ISSUE-01 to TC-ISSUE-10)", "134"),
        ("Table 4.3: Hall Manager Triage and Scoped Authorization Test Suite (TC-MGR-01 to TC-MGR-10)", "136"),
        ("Table 4.4: Maintenance Resolution and Work Order Test Suite (TC-MAINT-01 to TC-MAINT-08)", "138"),
        ("Table 4.5: Student Dispute and Auto-Closure Daemon Test Suite (TC-DISP-01 to TC-DISP-06)", "140"),
        ("Table 4.6: System Administration and Cross-Hall Security Test Suite (TC-ADMIN-01 to TC-ADMIN-06)", "141"),
        ("Table 4.7: API Benchmark Performance Metrics across Critical Endpoints (100 Concurrency)", "145"),
        ("Table 4.8: Empirical System Usability Scale (SUS) Evaluation Scores and Cohort Breakdown", "153")
    ]
    for tab_title, page_str in tables_list:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run(sanitize_text(tab_title))
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(11)
        dots_count = max(5, 75 - len(tab_title))
        r_dots = p.add_run(" " + ". " * (dots_count // 2) + " ")
        r_dots.font.name = 'Times New Roman'
        r_dots.font.size = Pt(10)
        r_dots.font.color.rgb = RGBColor(160, 160, 160)
        r2 = p.add_run(page_str)
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(11)
        r2.font.bold = True
    doc.add_page_break()

    # ==========================================
    # 9. LIST OF ABBREVIATIONS AND ACRONYMS
    # ==========================================
    print("Generating List of Abbreviations...")
    p_title_center("LIST OF ABBREVIATIONS AND ACRONYMS", size=14, bold=True, space_before=18, space_after=18)
    add_styled_table(
        "Table of Acronyms and Operational Abbreviations",
        ["Acronym", "Full Definition and Technical Context"],
        PRELIMS_DATA["abbreviations"],
        col_widths=[1.5, 4.5]
    )
    doc.add_page_break()

    # ==========================================
    # MAIN BODY: START OF SECTION 2 (ARABIC PAGE NUMBERING)
    # ==========================================
    section2 = doc.add_section()
    section2.top_margin = Inches(1.0)
    section2.bottom_margin = Inches(1.0)
    section2.left_margin = Inches(1.5)
    section2.right_margin = Inches(1.0)
    section2.header.is_linked_to_previous = False
    section2.footer.is_linked_to_previous = False

    # Footer with Arabic page numbers
    footer = section2.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer_p.add_run()
    footer_run.font.name = 'Times New Roman'
    footer_run.font.size = Pt(10)
    add_page_number(footer_run)

    # ==========================================
    # CHAPTER ONE: INTRODUCTION
    # ==========================================
    print("Generating Chapter One...")
    p_title_center(f"CHAPTER {CH1_DATA['chapter_number']}: {CH1_DATA['chapter_title']}", size=14, bold=True, space_before=18, space_after=18)
    for sec in CH1_DATA["sections"]:
        p_h1(f"{sec['num']} {sec['title']}")
        for p_text in sec["paragraphs"]:
            p_body(p_text)
    doc.add_page_break()

    # ==========================================
    # CHAPTER TWO: LITERATURE REVIEW
    # ==========================================
    print("Generating Chapter Two...")
    p_title_center(f"CHAPTER {CH2_DATA['chapter_number']}: {CH2_DATA['chapter_title']}", size=14, bold=True, space_before=18, space_after=18)
    for sec in CH2_DATA["sections"]:
        p_h1(f"{sec['num']} {sec['title']}")
        for p_text in sec["paragraphs"]:
            p_body(p_text)
        if "comparison_table" in sec:
            t_data = sec["comparison_table"]
            add_styled_table(t_data["caption"], t_data["headers"], t_data["rows"], col_widths=[1.2, 1.0, 1.0, 1.0, 1.0, 1.0])
    doc.add_page_break()

    # ==========================================
    # CHAPTER THREE: SYSTEM ANALYSIS AND DESIGN
    # ==========================================
    print("Generating Chapter Three...")
    p_title_center(f"CHAPTER {CH3_DATA['chapter_number']}: {CH3_DATA['chapter_title']}", size=14, bold=True, space_before=18, space_after=18)
    for sec in CH3_DATA["sections"]:
        p_h1(f"{sec['num']} {sec['title']}")
        for p_text in sec["paragraphs"]:
            p_body(p_text)

        if "functional_requirements_table" in sec:
            t = sec["functional_requirements_table"]
            add_styled_table(t["caption"], t["headers"], t["rows"], col_widths=[0.8, 1.1, 1.1, 2.2, 0.8])

        if "non_functional_requirements_table" in sec:
            t = sec["non_functional_requirements_table"]
            add_styled_table(t["caption"], t["headers"], t["rows"], col_widths=[0.8, 1.4, 2.3, 1.5])

        if "input_specifications_table" in sec:
            t = sec["input_specifications_table"]
            add_styled_table(t["caption"], t["headers"], t["rows"], col_widths=[1.2, 1.1, 0.8, 1.1, 1.8])

        if "data_dictionaries" in sec:
            for dd in sec["data_dictionaries"]:
                add_styled_table(dd["caption"], dd["headers"], dd["rows"], col_widths=[1.3, 1.1, 1.3, 2.3])

        if "figures" in sec:
            for fig in sec["figures"]:
                add_image_figure(fig["path"], fig["caption"])
                p_body(fig["description"], italic=True)

        if "use_case_specifications" in sec:
            for uc in sec["use_case_specifications"]:
                add_styled_table(uc["caption"], ["Specification Element", "Detailed Requirement Content"], uc["rows"], col_widths=[1.8, 4.2])

        if "algorithms" in sec:
            for algo in sec["algorithms"]:
                p_h2(algo["title"])
                add_code_block(algo["caption"], algo["code"])

        if "specs_table" in sec:
            t = sec["specs_table"]
            add_styled_table(t["caption"], t["headers"], t["rows"], col_widths=[1.2, 1.4, 1.7, 1.7])
    doc.add_page_break()

    # ==========================================
    # CHAPTER FOUR: IMPLEMENTATION AND EVALUATION
    # ==========================================
    print("Generating Chapter Four...")
    p_title_center(f"CHAPTER {CH4_DATA['chapter_number']}: {CH4_DATA['chapter_title']}", size=14, bold=True, space_before=18, space_after=18)
    for sec in CH4_DATA["sections"]:
        p_h1(f"{sec['num']} {sec['title']}")
        for p_text in sec["paragraphs"]:
            p_body(p_text)

        if "figures" in sec:
            for fig in sec["figures"]:
                add_image_figure(fig["path"], fig["caption"])
                p_body(fig["description"], italic=True)

        if "test_suite_tables" in sec:
            for tst in sec["test_suite_tables"]:
                add_styled_table(tst["caption"], tst["headers"], tst["rows"], col_widths=[1.0, 1.4, 1.4, 1.6, 0.6])

        if "performance_table" in sec:
            t = sec["performance_table"]
            add_styled_table(t["caption"], t["headers"], t["rows"], col_widths=[1.4, 0.8, 1.0, 0.9, 0.9, 0.9, 0.6])

        if "sus_table" in sec:
            t = sec["sus_table"]
            add_styled_table(t["caption"], t["headers"], t["rows"], col_widths=[1.5, 0.8, 1.1, 0.9, 0.8, 0.9])
    doc.add_page_break()

    # ==========================================
    # CHAPTER FIVE: CONCLUSION AND FUTURE WORKS
    # ==========================================
    print("Generating Chapter Five...")
    p_title_center(f"CHAPTER {CH5_DATA['chapter_number']}: {CH5_DATA['chapter_title']}", size=14, bold=True, space_before=18, space_after=18)
    for sec in CH5_DATA["sections"]:
        p_h1(f"{sec['num']} {sec['title']}")
        for p_text in sec["paragraphs"]:
            p_body(p_text)
    doc.add_page_break()

    # ==========================================
    # REFERENCES (APA 7th EDITION)
    # ==========================================
    print("Generating References...")
    p_title_center("REFERENCES", size=14, bold=True, space_before=18, space_after=18)
    for ref_text in REFERENCES_DATA:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
        r = p.add_run(sanitize_text(ref_text))
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
    doc.add_page_break()

    # ==========================================
    # APPENDICES
    # ==========================================
    print("Generating Appendices...")
    p_title_center("APPENDICES", size=14, bold=True, space_before=18, space_after=18)

    # APPENDIX A
    app_a = APPENDICES_DATA["appendix_a"]
    p_h1(app_a["title"])
    p_body(app_a["description"])
    for ep in app_a["endpoints"]:
        p_h2(f"{ep['method']} {ep['path']} ({ep['auth']})")
        p_body(ep["desc"], bold_prefix="Description: ")
        add_code_block(f"Payload Schema for {ep['method']} {ep['path']}", f"Request:\n{ep['request_body']}\n\nResponse:\n{ep['response']}")
    doc.add_page_break()

    # APPENDIX B
    app_b = APPENDICES_DATA["appendix_b"]
    p_h1(app_b["title"])
    p_body(app_b["description"])
    p_h2("B.1 User Entity Schema (User.ts)")
    add_code_block("Listing B.1: User.ts Mongoose Model Schema", app_b["user_model_code"])
    p_h2("B.2 Issue Entity Schema (Issue.ts)")
    add_code_block("Listing B.2: Issue.ts Mongoose Model Schema", app_b["issue_model_code"])
    if "issue_event_model_code" in app_b:
        p_h2("B.3 IssueEvent Audit Schema (IssueEvent.ts)")
        add_code_block("Listing B.3: IssueEvent.ts Mongoose Model Schema", app_b["issue_event_model_code"])
    if "assignment_model_code" in app_b:
        p_h2("B.4 Assignment Entity Schema (Assignment.ts)")
        add_code_block("Listing B.4: Assignment.ts Mongoose Model Schema", app_b["assignment_model_code"])
    doc.add_page_break()

    # APPENDIX C
    app_c = APPENDICES_DATA["appendix_c"]
    p_h1(app_c["title"])
    p_body(app_c["description"])
    add_code_block("Listing C.1: Supertest Integration Test Suite for Scoped Authorization", app_c["test_code"])
    doc.add_page_break()

    # APPENDIX D
    app_d = APPENDICES_DATA["appendix_d"]
    p_h1(app_d["title"])
    p_body(app_d["description"])
    p_h2("D.1 Production Environment Variables (.env)")
    add_code_block("Listing D.1: Production Environment Configuration", app_d["env_config"])
    p_h2("D.2 Docker Container Specification (Dockerfile)")
    add_code_block("Listing D.2: Multi-Stage Production Dockerfile", app_d["docker_config"])
    if "nginx_config" in app_d:
        p_h2("D.3 Nginx Reverse Proxy Configuration")
        add_code_block("Listing D.3: Nginx SSL Reverse Proxy Configuration", app_d["nginx_config"])
    doc.add_page_break()

    # APPENDIX E
    app_e = APPENDICES_DATA["appendix_e"]
    p_h1(app_e["title"])
    p_body(app_e["description"])
    p_body("The survey was administered using the standard 5-point Likert scale (1 = Strongly Disagree, 2 = Disagree, 3 = Neutral, 4 = Agree, 5 = Strongly Agree):")
    for q in app_e["survey_items"]:
        p_body(q, space_after=4)
    doc.add_page_break()

    # APPENDIX F
    app_f = APPENDICES_DATA["appendix_f"]
    p_h1(app_f["title"])
    p_body(app_f["description"])
    event_rows = [[e["code"], e["severity"], e["desc"]] for e in app_f["events"]]
    add_styled_table("Table F.1: Security and Lifecycle Event Catalog", ["Event Code", "Severity", "Description and Operational Trigger"], event_rows, col_widths=[2.0, 1.0, 3.0])

    print("Document compilation complete. Saving files...")
    out_docx = os.path.join(REPO_ROOT, "Gideon.docx")
    out_extless = os.path.join(REPO_ROOT, "Gideon")

    doc.save(out_docx)
    shutil.copyfile(out_docx, out_extless)
    print(f"Saved: {out_docx}")
    print(f"Saved: {out_extless}")

    # Validation checks
    total_words = sum(len(p.text.split()) for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                total_words += len(cell.text.split())

    em_dash_count = 0
    for p in doc.paragraphs:
        if '\u2014' in p.text:
            em_dash_count += 1
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                if '\u2014' in cell.text:
                    em_dash_count += 1

    print("=== VALIDATION SUMMARY ===")
    print(f"Paragraphs: {len(doc.paragraphs)}")
    print(f"Tables: {len(doc.tables)}")
    print(f"Inline Images / Shapes: {len(doc.inline_shapes)}")
    print(f"Total Words: {total_words}")
    print(f"Em dash count: {em_dash_count}")
    assert em_dash_count == 0, f"Error: {em_dash_count} em dashes found in document!"
    print("SUCCESS: 0 em dashes detected.")

if __name__ == "__main__":
    compile_document()
